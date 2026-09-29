"""PPTX rendering of ReportData with python-pptx (N-10; the `reports` extra).

A pure function of the verified snapshot: a title slide, the kind's sections in order (long lists and
tables continue on further slides), one slide per chart image from the shared PNG renderer, and a
provenance slide naming the snapshot hash that services/reports records. Nothing is recomputed.
"""
from __future__ import annotations

import io
from dataclasses import dataclass

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.presentation import Presentation as PptxPresentation
from pptx.util import Emu, Inches, Pt

from analystos.contracts.reports import ReportData
from analystos.reports import _common as C
from analystos.reports import _ooxml as O
from analystos.reports.charts import render_chart_png

ACCENT = RGBColor(47, 93, 138)
TEXT = RGBColor(31, 35, 40)
MUTED = RGBColor(89, 99, 110)
SEV = {"critical": RGBColor(207, 34, 46), "warning": RGBColor(154, 103, 0), "info": ACCENT}
LINES_PER_SLIDE = 12
ROWS_PER_SLIDE = 10
CHART_W, CHART_H = 1200, 560
W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.5)
TITLE_ONLY, BLANK = 5, 6


@dataclass
class Line:
    text: str
    level: int = 0
    bold: bool = False
    size: float = 16
    color: RGBColor = TEXT


def _slide(prs: PptxPresentation, title: str):
    s = prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY])
    t = s.shapes.title
    t.left, t.top, t.width, t.height = MARGIN, Inches(0.3), W - 2 * MARGIN, Inches(0.9)
    t.text = O.xml_text(title)
    for p in t.text_frame.paragraphs:
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb = Pt(26), True, ACCENT
    return s


def _body_box(s):
    box = s.shapes.add_textbox(MARGIN, Inches(1.3), W - 2 * MARGIN, H - Inches(1.8))
    box.text_frame.word_wrap = True
    return box.text_frame


def _lines(prs: PptxPresentation, title: str, lines: list[Line]) -> None:
    chunks = [lines[i : i + LINES_PER_SLIDE] for i in range(0, len(lines), LINES_PER_SLIDE)] or [[]]
    for n, chunk in enumerate(chunks):
        tf = _body_box(_slide(prs, title if n == 0 else f"{title} (continued)"))
        for k, ln in enumerate(chunk):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.level = ln.level
            r = p.add_run()
            r.text = O.xml_text(("• " if ln.level else "") + ln.text)
            r.font.size, r.font.bold, r.font.color.rgb = Pt(ln.size - 2 * ln.level), ln.bold, ln.color


def _table(prs: PptxPresentation, title: str, headers: list[str], rows: list[list[str]], note: str | None = None) -> None:
    for n in range(0, max(len(rows), 1), ROWS_PER_SLIDE):
        chunk = rows[n : n + ROWS_PER_SLIDE]
        s = _slide(prs, title if n == 0 else f"{title} (continued)")
        height = Emu(int(Inches(0.4)) * (len(chunk) + 1))
        tbl = s.shapes.add_table(len(chunk) + 1, len(headers), MARGIN, Inches(1.3), W - 2 * MARGIN, height).table
        for j, h in enumerate(headers):
            tbl.cell(0, j).text = O.xml_text(h)
        for i, row in enumerate(chunk, start=1):
            for j in range(len(headers)):
                tbl.cell(i, j).text = O.xml_text(row[j] if j < len(row) else "")
        for i in range(len(chunk) + 1):
            for j in range(len(headers)):
                for p in tbl.cell(i, j).text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(12 if i else 13)
        if note and n + ROWS_PER_SLIDE >= len(rows):
            box = s.shapes.add_textbox(MARGIN, H - Inches(0.8), W - 2 * MARGIN, Inches(0.4))
            r = box.text_frame.paragraphs[0].add_run()
            r.text, r.font.size, r.font.color.rgb = O.xml_text(note), Pt(11), MUTED


def _findings(d: ReportData) -> list[Line]:
    out = []
    for i in C.shown_findings(d):
        tag = f" [{i.change.upper()}]" if i.change else ""
        out.append(Line(f"{i.code}  {i.title}{tag}", bold=True, size=16))
        out.append(Line(O.finding_meta(i), size=11, color=MUTED))
        out.append(Line(i.finding, size=14))
        out.extend(Line(c, level=1, size=14) for c in i.caveats)
        if i.evidence_queries:
            out.append(Line("Evidence: " + ", ".join(q.id for q in i.evidence_queries), size=11, color=MUTED))
    return out


def _section(prs: PptxPresentation, key: str, d: ReportData) -> None:
    kind = d.kind
    title = C.section_title(key, kind)
    if key == "summary":
        lines = [] if d.summary_markdown.strip() else [Line(C.auto_summary(d))]
        for b in C.parse_markdown(d.summary_markdown):
            if b.kind == "para":
                lines.append(Line(C.inline_plain(b.items[0])))
            else:
                lines.extend(Line(C.inline_plain(it), level=1) for it in b.items)
        _lines(prs, title, lines)
    elif key == "kpis":
        if d.metrics:
            _table(prs, title, O.KPI_HEADERS, O.kpi_rows(d))
        else:
            _lines(prs, title, [Line("No metrics in this run.")])
    elif key == "findings":
        lines = _findings(d)
        _lines(prs, title, lines or [Line("No verified findings in this run." if kind == "executive" else "No findings in this run.")])
    elif key == "alerts":
        lines = []
        for a in C.sorted_alerts(d):
            lines.append(Line(f"{a.severity.upper()}  {a.title}", bold=True, color=SEV.get(a.severity, MUTED)))
            lines.append(Line(a.message + (f" (metric: {a.metric})" if a.metric else ""), size=14))
        _lines(prs, title, lines or [Line("No alerts.")])
    elif key == "changes":
        g = C.change_groups(d)
        lines = []
        if any(g.values()):
            for k, lab in (("new", "New"), ("changed", "Changed"), ("persisting", "Persisting"), ("resolved", "Resolved")):
                lines.append(Line(f"{lab} ({len(g[k])})", bold=True))
                lines.extend(Line(f"{i.code}  {i.title}", level=1) for i in g[k])
        _lines(prs, title, lines or [Line("No comparison with a previous run is available.")])
    elif key == "quality":
        if d.quality_issues:
            _table(prs, title, ["Severity", "Asset", "Column", "Issue"],
                   [[str(q.get("severity", "")), str(q.get("asset", "")), str(q.get("column", "") or ""), str(q.get("message", ""))]
                    for q in C.sorted_quality(d)])
        else:
            _lines(prs, title, [Line("No data quality issues recorded.")])
    elif key == "hypotheses":
        if d.hypotheses:
            _table(prs, title, ["Code", "Hypothesis", "Status", "Conclusion"],
                   [[str(h.get("code", "")), str(h.get("statement", "")), str(h.get("status", "")), str(h.get("conclusion", "") or "")]
                    for h in d.hypotheses])
        else:
            _lines(prs, title, [Line("No hypotheses were tested.")])
    elif key == "evidence":
        qs = [(i.code, q) for i in d.insights for q in i.evidence_queries]
        if qs:
            _table(prs, title, ["Finding", "Query", "Rows", "Result hash"], [[c, q.id, str(q.row_count), q.result_hash or ""] for c, q in qs])
        else:
            _lines(prs, title, [Line("No evidence queries recorded.")])
    elif key == "charts":
        if not d.charts:
            _lines(prs, title, [Line("No charts in this run.")])
        for ch in d.charts:
            rows = C.valid_rows(ch)
            if not rows:
                _lines(prs, f"{title}: {ch.title}", [Line("No data.", color=MUTED)])
                continue
            shown = rows[: C.MAX_TABLE_ROWS]
            _table(prs, f"{title}: {ch.title}", [str(c) for c in ch.columns], [[C.cell_text(v) for v in r] for r in shown],
                   note=f"Showing {len(shown)} of {len(rows)} rows." if len(rows) > len(shown) else None)


def _charts(prs: PptxPresentation, d: ReportData) -> None:
    for ch in d.charts:
        png = render_chart_png(ch, width_px=CHART_W, height_px=CHART_H)
        if not png:
            continue
        s = _slide(prs, ch.title)
        width = W - 2 * MARGIN
        s.shapes.add_picture(io.BytesIO(png), MARGIN, Inches(1.3), width=width, height=Emu(int(width * CHART_H / CHART_W)))


def build_pptx(data: ReportData) -> PptxPresentation:
    d = data
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    O.set_core(prs.core_properties, d)

    s = prs.slides.add_slide(prs.slide_layouts[BLANK])
    box = s.shapes.add_textbox(MARGIN, Inches(2.2), W - 2 * MARGIN, Inches(3.5))
    tf = box.text_frame
    tf.word_wrap = True
    head = [(C.KIND_LABELS.get(d.kind, d.kind).upper(), 14, True, ACCENT), (d.title, 36, True, TEXT),
            (f"{d.workspace_name}  |  generated {C.generated_label(d)}" + (f"  |  period {d.period}" if d.period else ""), 14, False, MUTED),
            (f"Objective: {d.objective}", 16, False, TEXT)]
    for k, (text, size, bold, color) in enumerate(head):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text, r.font.size, r.font.bold, r.font.color.rgb = O.xml_text(text), Pt(size), bold, color

    for key in C.sections_for(d.kind):
        _section(prs, key, d)
    _charts(prs, d)
    prov = [Line(f"Run: {d.run_id}", size=14), Line(f"Generated: {C.generated_label(d)}", size=14),
            Line(f"Report snapshot: {O.snapshot_hash(d)}", size=14)]
    if d.lineage_note:
        prov.append(Line(f"Lineage: {d.lineage_note}", size=14))
    prov.extend(Line(c, level=1, size=14) for c in d.caveats)
    prov.append(Line(C.CAUSATION_NOTE, size=12, color=MUTED))
    _lines(prs, C.SECTION_TITLES["provenance"], prov)
    return prs


def render_pptx(data: ReportData) -> bytes:
    buf = io.BytesIO()
    build_pptx(data).save(buf)
    return O.repack(buf.getvalue(), data)
