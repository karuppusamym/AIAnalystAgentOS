"""DOCX rendering of ReportData with python-docx (N-10; the `reports` extra).

A pure function of the verified snapshot, laid out like the PDF: the same sections per kind, the same
formatted values (reports._common), chart images from the same PNG renderer, and a provenance section
naming the snapshot hash that services/reports records. Nothing is recomputed from the source.
"""
from __future__ import annotations

import io

from docx import Document
from docx.document import Document as DocxDocument
from docx.shared import Mm, Pt, RGBColor

from analystos.contracts.reports import ReportData, ReportInsight
from analystos.reports import _common as C
from analystos.reports import _ooxml as O
from analystos.reports.charts import render_chart_png

ACCENT = RGBColor(47, 93, 138)
MUTED = RGBColor(89, 99, 110)
SEV = {"critical": RGBColor(207, 34, 46), "warning": RGBColor(154, 103, 0), "info": ACCENT}
CHART_W, CHART_H = 900, 420


def _p(doc: DocxDocument, text: str, *, size: float | None = None, bold: bool = False, italic: bool = False,
       color: RGBColor | None = None, style: str | None = None) -> None:
    run = doc.add_paragraph(style=style).add_run(O.xml_text(text))
    run.bold, run.italic = bold or None, italic or None
    if size:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color


def _bullets(doc: DocxDocument, items: list[str]) -> None:
    for it in items:
        doc.add_paragraph(O.xml_text(it), style="List Bullet")


def _table(doc: DocxDocument, headers: list[str], rows: list[list[str]]) -> None:
    if not rows:
        return
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for cell, h in zip(t.rows[0].cells, headers, strict=True):
        cell.text = ""
        cell.paragraphs[0].add_run(O.xml_text(h)).bold = True
    for row in rows:
        cells = t.add_row().cells
        for cell, v in zip(cells, row, strict=False):
            cell.text = O.xml_text(v)
    doc.add_paragraph()


def _insight(doc: DocxDocument, i: ReportInsight, with_evidence: bool) -> None:
    tag = f" [{i.change.upper()}]" if i.change else ""
    doc.add_heading(O.xml_text(f"{i.code}  {i.title}{tag}"), level=3)
    _p(doc, O.finding_meta(i), size=8, color=MUTED)
    _p(doc, i.finding)
    if i.caveats:
        _p(doc, "Caveats:", bold=True, size=9)
        _bullets(doc, i.caveats)
    if not i.evidence_queries:
        return
    if with_evidence:
        for q in i.evidence_queries:
            h = f", hash {q.result_hash}" if q.result_hash else ""
            _p(doc, f"Evidence query {q.id} ({q.row_count} rows{h})", size=8, color=MUTED)
            run = doc.add_paragraph().add_run(O.xml_text(q.sql.rstrip()))
            run.font.name, run.font.size = "Consolas", Pt(8)
    else:
        _p(doc, "Evidence: " + ", ".join(q.id for q in i.evidence_queries), size=8, color=MUTED)


def _section(doc: DocxDocument, key: str, d: ReportData) -> None:
    kind = d.kind
    doc.add_heading(C.section_title(key, kind), level=1)
    if key == "summary":
        if not d.summary_markdown.strip():
            _p(doc, C.auto_summary(d))
        for b in C.parse_markdown(d.summary_markdown):
            if b.kind == "para":
                _p(doc, C.inline_plain(b.items[0]))
            else:
                _bullets(doc, [C.inline_plain(it) for it in b.items])
    elif key == "kpis":
        if not d.metrics:
            _p(doc, "No metrics in this run.")
        _table(doc, O.KPI_HEADERS, O.kpi_rows(d))
    elif key == "findings":
        items = C.shown_findings(d)
        if not items:
            _p(doc, "No verified findings in this run." if kind == "executive" else "No findings in this run.")
        for i in items:
            _insight(doc, i, kind == "statistical")
    elif key == "alerts":
        if not d.alerts:
            _p(doc, "No alerts.")
        for a in C.sorted_alerts(d):
            _p(doc, f"{a.severity.upper()}  {a.title}", bold=True, color=SEV.get(a.severity, MUTED))
            _p(doc, a.message + (f" (metric: {a.metric})" if a.metric else ""))
    elif key == "changes":
        g = C.change_groups(d)
        if not any(g.values()):
            _p(doc, "No comparison with a previous run is available.")
            return
        for k, lab in (("new", "New"), ("changed", "Changed"), ("persisting", "Persisting"), ("resolved", "Resolved")):
            _p(doc, f"{lab} ({len(g[k])})", bold=True)
            _bullets(doc, [f"{i.code}  {i.title}" for i in g[k]] or ["none"])
    elif key == "quality":
        if not d.quality_issues:
            _p(doc, "No data quality issues recorded.")
        _table(doc, ["Severity", "Asset", "Column", "Issue"],
               [[str(q.get("severity", "")), str(q.get("asset", "")), str(q.get("column", "") or ""), str(q.get("message", ""))]
                for q in C.sorted_quality(d)])
    elif key == "hypotheses":
        if not d.hypotheses:
            _p(doc, "No hypotheses were tested.")
        _table(doc, ["Code", "Hypothesis", "Status", "Conclusion"],
               [[str(h.get("code", "")), str(h.get("statement", "")), str(h.get("status", "")), str(h.get("conclusion", "") or "")]
                for h in d.hypotheses])
    elif key == "evidence":
        qs = [(i.code, q) for i in d.insights for q in i.evidence_queries]
        if not qs:
            _p(doc, "No evidence queries recorded.")
        _table(doc, ["Finding", "Query", "Rows", "Result hash"], [[c, q.id, str(q.row_count), q.result_hash or ""] for c, q in qs])
    elif key == "charts":
        if not d.charts:
            _p(doc, "No charts in this run.")
        for ch in d.charts:
            rows = C.valid_rows(ch)
            doc.add_heading(O.xml_text(ch.title), level=3)
            if not rows:
                _p(doc, "No data.", size=8.5, color=MUTED)
                continue
            _table(doc, [str(c) for c in ch.columns], [[C.cell_text(v) for v in r] for r in rows[: C.MAX_TABLE_ROWS]])
            if len(rows) > C.MAX_TABLE_ROWS:
                _p(doc, f"Showing {C.MAX_TABLE_ROWS} of {len(rows)} rows.", size=8, color=MUTED)


def _charts(doc: DocxDocument, d: ReportData, width_mm: float) -> None:
    images = [(ch, png) for ch in d.charts if (png := render_chart_png(ch, width_px=CHART_W, height_px=CHART_H))]
    if not images:
        return
    doc.add_heading("Charts", level=1)
    for _, png in images:
        doc.add_picture(io.BytesIO(png), width=Mm(width_mm))


def _provenance(doc: DocxDocument, d: ReportData) -> None:
    doc.add_heading(C.SECTION_TITLES["provenance"], level=1)
    _p(doc, f"Run: {d.run_id}", size=9)
    _p(doc, f"Generated: {C.generated_label(d)}", size=9)
    _p(doc, f"Report snapshot: {O.snapshot_hash(d)}", size=9)
    if d.lineage_note:
        _p(doc, f"Lineage: {d.lineage_note}", size=9)
    if d.caveats:
        _p(doc, "Caveats:", size=9, bold=True)
        _bullets(doc, d.caveats)
    _p(doc, C.CAUSATION_NOTE, size=9, italic=True)


def build_docx(data: ReportData) -> DocxDocument:
    d = data
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Mm(16)
    doc.styles["Normal"].font.size = Pt(10)
    O.set_core(doc.core_properties, d)
    sec.footer.paragraphs[0].text = O.xml_text(f"{d.title} | run {d.run_id} | generated {C.generated_label(d)}")

    _p(doc, C.KIND_LABELS.get(d.kind, d.kind).upper(), size=8.5, bold=True, color=ACCENT)
    doc.add_heading(O.xml_text(d.title), level=0)
    _p(doc, f"{d.workspace_name}  |  generated {C.generated_label(d)}" + (f"  |  period {d.period}" if d.period else ""),
       size=9, color=MUTED)
    _p(doc, f"Objective: {d.objective}")
    for key in C.sections_for(d.kind):
        _section(doc, key, d)
    _charts(doc, d, 178)
    _provenance(doc, d)
    return doc


def render_docx(data: ReportData) -> bytes:
    buf = io.BytesIO()
    build_docx(data).save(buf)
    return O.repack(buf.getvalue(), data)
