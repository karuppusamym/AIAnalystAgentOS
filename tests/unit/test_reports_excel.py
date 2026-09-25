"""XLSX report rendering: sheets, formats, formula-injection safety, native charts, determinism."""
from __future__ import annotations

import io
import zipfile

from openpyxl import load_workbook
from tests.report_sample import sample_report

from analystos.contracts.reports import ReportChart, ReportData
from analystos.reports import render
from analystos.reports.excel import render_xlsx, sheet_title

FIXED = ["Summary", "KPIs", "Findings", "Evidence", "Quality", "Hypotheses", "Alerts"]


def _wb(d: ReportData):
    return load_workbook(io.BytesIO(render_xlsx(d)))


def _values(wb) -> dict[str, list[tuple]]:
    return {ws.title: [tuple(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def test_sheets_and_headers():
    wb = _wb(sample_report())
    names = wb.sheetnames
    assert names[:4] == ["Summary", "KPIs", "Findings", "Evidence"] and names[-3:] == ["Quality", "Hypotheses", "Alerts"]
    chart_sheets = names[4:-3]
    assert len(chart_sheets) == 5 and len(set(n.lower() for n in chart_sheets)) == 5
    assert all(len(n) <= 31 and not set(n) & set("[]:*?/\\") for n in chart_sheets)
    assert chart_sheets[1].startswith("Median resolution by group _hou")
    for name in FIXED:
        ws = wb[name]
        if name != "Summary":
            assert ws.freeze_panes == "A2"
        assert ws["A1"].font.b
    assert [c.value for c in wb["KPIs"][1]] == ["Metric", "Display name", "Definition", "Value", "Previous", "Change", "Format"]
    s = {r[0]: r[1] for r in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    assert s["Run id"] == "run_0f3a9c2e" and s["Generated at"] == "2026-09-21T08:30:00Z"
    assert "SLA attainment fell 1.8 pp" in s["Summary"]


def test_kpi_formats():
    ws = _wb(sample_report())["KPIs"]
    rows = {r[0].value: r for r in ws.iter_rows(min_row=2)}
    sla = rows["sla_rate"]
    assert sla[3].value == 0.873 and sla[3].number_format == "0.0%"
    assert sla[4].value == 0.891 and sla[4].number_format == "0.0%"
    assert abs(sla[5].value - (-0.018)) < 1e-9 and "%" in sla[5].number_format
    assert sla[6].value == "percent"
    assert rows["mttr"][3].number_format == '#,##0.0" h"'
    assert rows["first_response"][4].value is None and rows["first_response"][5].value is None
    assert rows["backlog"][3].value == "n/a" and rows["backlog"][4].value == 212


def test_formula_injection_is_stored_as_text():
    raw = render_xlsx(sample_report())
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        for n in z.namelist():
            if n.startswith("xl/worksheets/sheet"):
                assert b"<f>" not in z.read(n), n  # no formula cells anywhere
    wb = load_workbook(io.BytesIO(raw))
    f = {r[0].value: r for r in wb["Findings"].iter_rows(min_row=2)}["F-003"]
    assert f[1].value.startswith("=HYPERLINK") and f[1].data_type == "s"
    assert f[2].value.startswith("+cmd") and f[2].data_type == "s"
    assert f[5].value.startswith("@SUM") and f[5].data_type == "s"
    q = [r for r in wb["Quality"].iter_rows(min_row=2) if str(r[3].value).startswith("=1+1")]
    assert q and q[0][3].data_type == "s"
    k = {r[0].value: r for r in wb["KPIs"].iter_rows(min_row=2)}["backlog"]
    assert k[2].value.startswith("-SUM") and k[2].data_type == "s"


def test_findings_evidence_alerts_content():
    wb = _wb(sample_report())
    f = {r[0]: r for r in wb["Findings"].iter_rows(min_row=2, values_only=True)}
    assert f["F-001"][3] == 0.92 and f["F-001"][4] == "yes" and f["F-001"][6] == "persisting"
    assert f["F-002"][5] == "Holiday week — lower staffing…; Zürich site data is partial (日本語 labels)."
    assert f["F-000"][6] == "resolved"
    assert wb["Findings"]["D2"].number_format == "0%"
    ev = list(wb["Evidence"].iter_rows(min_row=2, values_only=True))
    assert ev[0][:4] == ("F-001", "q-sla-by-priority", 4, "sha256:9f2c1a") and ev[0][4].startswith("SELECT priority")
    assert [r[0] for r in wb["Alerts"].iter_rows(min_row=2, values_only=True)] == ["critical", "warning"]
    assert len(list(wb["Hypotheses"].iter_rows(min_row=2))) == 3


def test_chart_sheets_have_data_and_native_charts():
    wb = _wb(sample_report())
    names = wb.sheetnames[4:-3]
    line, bar, pie, empty, broken = (wb[n] for n in names)
    assert [c.value for c in line[1]][:2] == ["week", "sla_rate"] and line["B2"].value == 0.88
    assert bar["A5"].value == "=cmd()" and bar["A5"].data_type == "s"
    for ws in (line, bar, pie):
        assert len(ws._charts) == 1
    assert empty["A2"].value == "No data." and not empty._charts
    assert not broken._charts
    kinds = [type(wb[n]._charts[0]).__name__ for n in names[:3]]
    assert kinds == ["LineChart", "BarChart", "PieChart"]


def test_native_chart_types_and_pivot():
    d = ReportData(title="t", workspace_name="w", objective="o", run_id="r", generated_at="2026-01-01T00:00:00Z", charts=[
        ReportChart(key="s", title="Scatter", chart_type="scatter", columns=["x", "y"], rows=[[1, 2], [2, 3.5], [3, 1]]),
        ReportChart(key="sb", title="Stacked", chart_type="stacked_bar", columns=["wk", "prio", "n"],
                    rows=[["W1", "P1", 1], ["W1", "P2", 2], ["W2", "P1", 3]]),
        ReportChart(key="t", title="Scatter", chart_type="table", columns=["a"], rows=[["x"]]),
    ])
    wb = _wb(d)
    s, sb, t = (wb[n] for n in wb.sheetnames[4:-3])
    assert wb.sheetnames[4:-3] == ["Scatter", "Stacked", "Scatter (2)"]
    assert type(s._charts[0]).__name__ == "ScatterChart"
    ch = sb._charts[0]
    assert type(ch).__name__ == "BarChart" and ch.grouping == "stacked"
    assert [c.value for c in sb[1]][4:7] == ["wk", "P1", "P2"]  # pivoted block beside the raw rows
    assert sb["F3"].value == 3 and sb["G3"].value is None
    assert not t._charts


def test_sheet_title_sanitizing():
    taken = {"summary", "kpis"}
    assert sheet_title("Summary", taken) == "Summary (2)"
    assert sheet_title("a/b:c*d?e[f]g\\h", taken) == "a_b_c_d_e_f_g_h"
    long = sheet_title("x" * 40, taken)
    assert len(long) == 31 and len(sheet_title("x" * 40, taken)) == 31 and sheet_title("x" * 40, taken).endswith("(3)")
    assert sheet_title("''", taken) == "Chart"
    assert sheet_title("History", taken) == "History (2)"


def test_xlsx_content_is_deterministic_and_dispatch():
    assert _values(_wb(sample_report("operational"))) == _values(_wb(sample_report("operational")))
    body, mime, ext = render(sample_report(), "xlsx")
    assert ext == "xlsx" and mime.endswith("spreadsheetml.sheet") and body[:2] == b"PK"
    wb = load_workbook(io.BytesIO(body))
    assert wb.properties.created.isoformat().startswith("2026-09-21T08:30")
