"""N-10: PPTX and DOCX reports are rendered from the verified ReportData snapshot only. Reopened, every
number and finding is the snapshot's own (formatted by reports._common, never recomputed), the snapshot
hash is in the document and on the artifact, and the same snapshot yields the same bytes."""
from __future__ import annotations

import io
import zipfile
from datetime import UTC, datetime
from functools import cache
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation
from tests.report_sample import sample_report
from tests.unit.step_fixtures import WS, world  # noqa: F401

from analystos.contracts.reports import ReportData, ReportInsight, ReportMetric
from analystos.core import profiles
from analystos.core.errors import FeatureUnavailable
from analystos.core.ids import stable_hash
from analystos.reports import _common as C
from analystos.reports import _ooxml as O
from analystos.reports import render, unavailable_formats

KINDS = ["executive", "operational", "statistical", "exception", "weekly_summary"]
MIME = {"pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}


@cache
def _render(kind: str, fmt: str) -> bytes:
    body, mime, ext = render(sample_report(kind), fmt)
    assert mime == MIME[fmt] and ext == fmt
    return body


def _docx(b: bytes) -> tuple[str, list[list[list[str]]], object]:
    doc = Document(io.BytesIO(b))
    tables = [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]
    text = "\n".join([p.text for p in doc.paragraphs] + [c for t in tables for r in t for c in r])
    return text, tables, doc


def _pptx(b: bytes) -> tuple[str, list[list[list[str]]], object]:
    prs = Presentation(io.BytesIO(b))
    texts, tables = [], []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                texts.append(sh.text_frame.text)
            if sh.has_table:
                tables.append([[c.text for c in r.cells] for r in sh.table.rows])
    return "\n".join(texts + [c for t in tables for r in t for c in r]), tables, prs


READ = {"docx": _docx, "pptx": _pptx}


def _kpi_rows(tables: list[list[list[str]]]) -> list[list[str]]:
    """KPI table rows, joined across continuation slides."""
    return [r for t in tables if t[0] == O.KPI_HEADERS for r in t[1:]]


@pytest.mark.parametrize("fmt", ["docx", "pptx"])
@pytest.mark.parametrize("kind", KINDS)
def test_office_report_reopens_with_the_snapshot_numbers_and_hash(kind, fmt):
    d = sample_report(kind)
    text, tables, obj = READ[fmt](_render(kind, fmt))
    assert _kpi_rows(tables) == O.kpi_rows(d)  # every KPI cell is the snapshot value, formatted like the PDF
    for m in d.metrics:
        assert C.format_value(m.value, m.format) in text
    for i in C.shown_findings(d):
        assert O.xml_text(i.finding) in text and i.code in text
    digest = stable_hash(d.model_dump())
    assert f"Report snapshot: {digest}" in text and obj.core_properties.identifier == f"analystos-report-data:{digest}"
    assert d.run_id in text and C.CAUSATION_NOTE in text
    assert "日本語" in text  # unicode kept as written
    assert obj.core_properties.created.isoformat().startswith("2026-09-21T08:30")  # generated_at, in UTC


def test_sections_follow_the_kind():
    ex, _, _ = _docx(_render("executive", "docx"))
    assert "Top verified findings" in ex and "Data quality issues" not in ex
    st, st_tables, _ = _docx(_render("statistical", "docx"))
    assert st.index("Hypotheses and conclusions") < st.index("Findings") and "SELECT priority, AVG" in st
    assert any(t[0] == ["Finding", "Query", "Rows", "Result hash"] for t in st_tables)
    titles = [s.shapes.title.text for s in _pptx(_render("exception", "pptx"))[2].slides if s.shapes.title is not None]
    assert titles.index("Alerts") < titles.index("Summary") and titles[-1] == "Provenance"


def test_charts_are_embedded_from_the_shared_png_renderer():
    _, _, doc = _docx(_render("operational", "docx"))
    assert len(doc.inline_shapes) >= 1
    _, _, prs = _pptx(_render("operational", "pptx"))
    assert sum(1 for s in prs.slides for sh in s.shapes if sh.shape_type == 13) >= 1  # MSO_SHAPE_TYPE.PICTURE


@pytest.mark.parametrize("fmt", ["docx", "pptx"])
def test_bytes_are_deterministic_and_zip_entries_carry_generated_at(fmt):
    again, _, _ = render(sample_report("operational"), fmt)
    assert again == _render("operational", fmt)
    with zipfile.ZipFile(io.BytesIO(again)) as z:
        assert {i.date_time for i in z.infolist()} == {(2026, 9, 21, 8, 30, 0)}


@pytest.mark.parametrize("fmt", ["docx", "pptx"])
def test_hostile_text_and_many_rows_do_not_break_the_document(fmt):
    metrics = [ReportMetric(name=f"m{n}", display_name=f"Metric {n}\x07", value=n * 1.5, previous_value=n, format="hours")
               for n in range(25)]
    d = ReportData(title="Отчёт \x00 報告 🚀", workspace_name="ws\x0b", objective="Ω\ud800", run_id="r",
                   generated_at="1970-01-01T00:00:00Z", metrics=metrics,
                   insights=[ReportInsight(code="Ж", title="χ²", finding="√ 😀 \x1f", confidence=0.5, verified=True)])
    text, tables, _ = READ[fmt](render(d, fmt)[0])
    assert "Отчёт  報告 🚀" in text and "√ 😀 " in text
    assert _kpi_rows(tables) == [[O.xml_text(c) for c in r] for r in O.kpi_rows(d)]  # 25 rows, over several slides


def test_office_formats_say_why_they_are_unavailable(monkeypatch):
    monkeypatch.setattr(profiles, "_importable", lambda m: m != "pptx")
    assert set(unavailable_formats(("html", "pdf", "pptx", "docx"))) == {"pptx"}
    assert "python-pptx" in unavailable_formats(("pptx",))["pptx"]
    with pytest.raises(FeatureUnavailable, match="PPTX reports are unavailable"):
        render(sample_report(), "pptx")


def test_generate_report_records_the_snapshot_hash_for_office_files(world, monkeypatch, tmp_path):  # noqa: F811
    from analystos.artifacts.registry import save_artifact
    from analystos.core.config import get_settings
    from analystos.db.base import session_scope
    from analystos.db.models import AnalysisRun
    from analystos.services import reports as svc

    monkeypatch.setattr(get_settings(), "artifact_dir", tmp_path)
    with session_scope() as s:
        s.add(AnalysisRun(id="run_n10", workspace_id=WS, objective="Why are orders late?", status="COMPLETED",
                          requested_by="usr_owner", finished_at=datetime(2026, 9, 21, 8, 30, tzinfo=UTC),
                          summary={"summary_markdown": "Late orders rose to **12%**."}))
    with session_scope() as s:
        save_artifact(s, workspace_id=WS, run_id="run_n10", type_="metric", name="late_rate", status="final",
                      content={"display_name": "Late rate", "format": "percent", "validation": {"value": 0.1234}})
    with session_scope() as s:
        snapshot = svc.build_report_data(s, "run_n10", "executive")
        art = svc.generate_report(s, "run_n10", kind="executive", formats=("pptx", "docx"), actor="user:usr_owner")
        digest = stable_hash(snapshot.model_dump())
        assert art.content["report_data_hash"] == digest and set(art.content["files"]) == {"pptx", "docx"}
        for fmt in ("pptx", "docx"):
            body, mime, ext = svc.report_file(art, fmt)
            assert mime == MIME[fmt] and ext == fmt and Path(art.content["files"][fmt]["path"]).name.endswith(f"{digest[:12]}.{fmt}")
            assert body == render(snapshot, fmt)[0]  # the stored file is exactly the snapshot's rendering
            text, tables, obj = READ[fmt](body)
            assert obj.core_properties.identifier == f"analystos-report-data:{digest}"
            assert _kpi_rows(tables) == [["Late rate", "12.3%", "n/a", "n/a", ""]]
            assert "Late orders rose to 12%." in text


def test_the_formats_are_offered_by_the_service_and_api():
    from analystos.reports import FORMATS
    from analystos.services.reports import FORMATS as SERVICE_FORMATS

    assert {"pptx", "docx"} <= set(FORMATS) and {"pptx", "docx"} <= set(SERVICE_FORMATS)
